#!/usr/bin/env python3
"""Convert pilot XLSX files into a stable, pseudonymized CrewQual CSV draft.

The converter deliberately refuses ambiguous identity matches. It never writes a
name mapping and never accepts the HMAC key through a command-line argument.
"""

from __future__ import annotations

import argparse
import base64
import calendar
import colorsys
import csv
import hashlib
import hmac
import io
import os
import re
import sys
import unicodedata
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Callable, Iterable, Sequence


KEY_ENV_NAME = "CREWQUAL_MOCK_HMAC_KEY_HEX"
DEFAULT_SOURCE = "测试.xlsx"
DEFAULT_TEMPLATE = "CrewQual-飞行员批量导入模板-当前.csv"
DEFAULT_OUTPUT = "CrewQual-飞行员批量导入-待补.csv"

EXPECTED_HEADERS = [
    "employeeNumber",
    "displayName",
    "mobile",
    "aircraftType",
    "roleCode",
    "unitCode",
    "rankCode",
    "medical-certificate.issueDate",
    "medical-certificate.trainingDate",
    "medical-certificate.expiryDate",
    "medical-certificate.levelOrParameter",
    "annual-recurrent-training.issueDate",
    "annual-recurrent-training.trainingDate",
    "annual-recurrent-training.expiryDate",
    "annual-recurrent-training.levelOrParameter",
    "dangerous-goods-training.issueDate",
    "dangerous-goods-training.trainingDate",
    "dangerous-goods-training.expiryDate",
    "dangerous-goods-training.levelOrParameter",
    "icao-english-endorsement.issueDate",
    "icao-english-endorsement.trainingDate",
    "icao-english-endorsement.expiryDate",
    "icao-english-endorsement.levelOrParameter",
    "chinese-language-assessment.issueDate",
    "chinese-language-assessment.trainingDate",
    "chinese-language-assessment.expiryDate",
    "chinese-language-assessment.levelOrParameter",
    "simulator-recurrent-training.issueDate",
    "simulator-recurrent-training.trainingDate",
    "simulator-recurrent-training.expiryDate",
    "simulator-recurrent-training.levelOrParameter",
]

QUALIFICATION_GROUPS = {
    "medical-certificate": False,
    "annual-recurrent-training": True,
    "dangerous-goods-training": False,
    "icao-english-endorsement": False,
    "chinese-language-assessment": False,
    "simulator-recurrent-training": True,
}


class ConversionError(RuntimeError):
    """An expected, privacy-safe error that can be shown to the operator."""


@dataclass(frozen=True)
class HeaderLocation:
    sheet_name: str
    row: int
    columns: dict[str, int]


@dataclass(frozen=True)
class SourceRow:
    row_number: int
    normalized_name: str
    raw_values: dict[str, Any]
    raw_cells: dict[str, Any]


@dataclass(frozen=True)
class PersonRow:
    row_number: int
    person_key: str
    token: str
    normalized_name: str
    real_phone: str
    source: SourceRow


@dataclass(frozen=True)
class ConversionSummary:
    source_rows: int
    matched_phones: int
    qualification_counts: dict[str, int]
    suppressed_qualifications: int


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return unicodedata.normalize("NFKC", str(value)).strip()


def compact_header(value: Any) -> str:
    text = normalize_text(value)
    return re.sub(r"[\s\"“”]+", "", text).casefold()


def normalize_name(value: Any) -> str:
    text = normalize_text(value).casefold()
    return "".join(character for character in text if not character.isspace())


def normalize_phone(value: Any) -> str | None:
    text = normalize_text(value)
    if not text:
        return None
    candidates = re.findall(r"(?<!\d)(?:\+?86[-\s]?)?(1\d{10})(?!\d)", text)
    distinct = list(dict.fromkeys(candidates))
    if len(distinct) == 1:
        return distinct[0]
    digits = re.sub(r"\D", "", text)
    if len(digits) == 13 and digits.startswith("86"):
        digits = digits[2:]
    return digits if len(digits) == 11 else None


def load_hmac_key(environ: dict[str, str] | None = None) -> bytes:
    raw = (environ or os.environ).get(KEY_ENV_NAME, "").strip()
    if not re.fullmatch(r"[0-9a-fA-F]{64}", raw):
        raise ConversionError(
            f"环境变量 {KEY_ENV_NAME} 必须是 64 个十六进制字符（32 字节），且不得作为命令行参数传入。"
        )
    return bytes.fromhex(raw)


def hmac_digest(key: bytes, purpose: str, person_key: str, counter: int | None = None) -> bytes:
    domain = purpose if counter is None else f"{purpose}:{counter}"
    message = f"{domain}\x1f{person_key}".encode("utf-8")
    return hmac.new(key, message, hashlib.sha256).digest()


def person_token(key: bytes, person_key: str) -> str:
    encoded = base64.b32encode(hmac_digest(key, "person-id", person_key)).decode("ascii")
    return encoded.rstrip("=")[:16]


def assign_mock_mobiles(key: bytes, people: Sequence[PersonRow]) -> dict[str, str]:
    used: set[str] = set()
    assigned: dict[str, str] = {}
    for person in sorted(people, key=lambda item: item.token):
        for counter in range(1_000):
            digest = hmac_digest(key, "mobile", person.person_key, counter)
            suffix = int.from_bytes(digest[:8], "big") % 100_000_000
            candidate = f"100{suffix:08d}"
            if candidate not in used:
                used.add(candidate)
                assigned[person.token] = candidate
                break
        else:
            raise ConversionError("无法生成唯一 mock 手机号；请更换 HMAC 密钥后重试。")
    return assigned


def add_months(value: date, months: int) -> date:
    target_index = value.year * 12 + (value.month - 1) + months
    target_year, target_month_zero = divmod(target_index, 12)
    target_month = target_month_zero + 1
    target_day = min(value.day, calendar.monthrange(target_year, target_month)[1])
    return date(target_year, target_month, target_day)


def iso_date(value: date) -> str:
    return value.isoformat()


def is_blank(value: Any) -> bool:
    return value is None or normalize_text(value) == ""


def parse_date_value(value: Any, excel_epoch: Any = None) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        if not 1 <= float(value) <= 100_000:
            return None
        try:
            from openpyxl.utils.datetime import from_excel

            parsed = from_excel(value, epoch=excel_epoch)
            return parsed.date() if isinstance(parsed, datetime) else parsed
        except (ImportError, OverflowError, TypeError, ValueError):
            return None
    text = normalize_text(value)
    if not text or text[0] in "=+@":
        return None
    patterns = (
        r"(?P<year>\d{4})[-/.](?P<month>\d{1,2})[-/.](?P<day>\d{1,2})",
        r"(?P<year>\d{4})年(?P<month>\d{1,2})月(?P<day>\d{1,2})日?",
    )
    for pattern in patterns:
        match = re.fullmatch(pattern, text)
        if match:
            try:
                return date(
                    int(match.group("year")),
                    int(match.group("month")),
                    int(match.group("day")),
                )
            except ValueError:
                return None
    return None


def parse_level(value: Any, raw_cell: Any = None) -> str | None:
    if raw_cell is not None and getattr(raw_cell, "data_type", None) == "f":
        return None
    if is_blank(value):
        return None
    text = normalize_text(value)
    if text[0] in "=+-@" or any(ord(character) < 32 for character in text):
        return None
    return text if len(text) <= 128 else None


def parse_cell_date(value: Any, raw_cell: Any, excel_epoch: Any) -> tuple[date | None, bool]:
    if raw_cell is not None and getattr(raw_cell, "data_type", None) == "f":
        return None, True
    if is_blank(value):
        return None, False
    parsed = parse_date_value(value, excel_epoch)
    return parsed, parsed is None


def theme_palette(workbook: Any) -> list[str | None]:
    raw_theme = getattr(workbook, "loaded_theme", None)
    if not raw_theme:
        return []
    try:
        root = ET.fromstring(raw_theme)
        scheme = next(element for element in root.iter() if element.tag.endswith("clrScheme"))
    except (ET.ParseError, StopIteration, TypeError):
        return []
    colors_by_name: dict[str, str | None] = {}
    for item in list(scheme):
        color_value: str | None = None
        for color in list(item):
            color_kind = color.tag.rsplit("}", 1)[-1]
            color_value = (
                color.attrib.get("lastClr") or color.attrib.get("val")
                if color_kind == "sysClr"
                else color.attrib.get("val") or color.attrib.get("lastClr")
            )
            if color_value:
                break
        colors_by_name[item.tag.rsplit("}", 1)[-1]] = color_value
    # OOXML theme indices use this order even though clrScheme XML commonly
    # serializes dk1 before lt1.
    order = (
        "lt1",
        "dk1",
        "lt2",
        "dk2",
        "accent1",
        "accent2",
        "accent3",
        "accent4",
        "accent5",
        "accent6",
        "hlink",
        "folHlink",
    )
    return [colors_by_name.get(name) for name in order]


def apply_tint(rgb: tuple[int, int, int], tint: float) -> tuple[int, int, int]:
    if not tint:
        return rgb
    if tint < 0:
        return tuple(max(0, min(255, round(channel * (1 + tint)))) for channel in rgb)
    return tuple(
        max(0, min(255, round(channel * (1 - tint) + 255 * tint))) for channel in rgb
    )


def color_to_rgb(color: Any, workbook: Any) -> tuple[int, int, int] | None:
    if color is None:
        return None
    raw: str | None = None
    color_type = getattr(color, "type", None)
    if color_type == "rgb":
        raw = getattr(color, "rgb", None)
    elif color_type == "theme":
        palette = theme_palette(workbook)
        theme_index = getattr(color, "theme", None)
        if isinstance(theme_index, int) and 0 <= theme_index < len(palette):
            raw = palette[theme_index]
    elif color_type == "indexed":
        try:
            from openpyxl.styles.colors import COLOR_INDEXED

            index = int(getattr(color, "indexed"))
            raw = COLOR_INDEXED[index] if 0 <= index < len(COLOR_INDEXED) else None
        except (ImportError, TypeError, ValueError):
            raw = None
    if not raw:
        return None
    raw = str(raw).strip().lstrip("#")[-6:]
    if not re.fullmatch(r"[0-9A-Fa-f]{6}", raw):
        return None
    rgb = tuple(int(raw[index : index + 2], 16) for index in (0, 2, 4))
    return apply_tint(rgb, float(getattr(color, "tint", 0.0) or 0.0))


def rgb_is_blue(rgb: tuple[int, int, int]) -> bool:
    red, green, blue = rgb
    hue, saturation, value = colorsys.rgb_to_hsv(red / 255, green / 255, blue / 255)
    return (
        0.50 <= hue <= 0.75
        and saturation >= 0.08
        and value >= 0.20
        and blue >= red + 10
        and blue >= green + 5
    )


def medical_cycle_months(raw_cell: Any, workbook: Any) -> int | None:
    fill = getattr(raw_cell, "fill", None)
    fill_type = getattr(fill, "fill_type", None)
    colors_to_check: list[Any] = []
    unresolved_colored_style = False
    if fill_type and fill_type != "none":
        colors_to_check.append(getattr(fill, "fgColor", None))
    font_color = getattr(getattr(raw_cell, "font", None), "color", None)
    if font_color is not None and getattr(font_color, "type", None) in {
        "rgb",
        "theme",
        "indexed",
    }:
        colors_to_check.append(font_color)
    for color in colors_to_check:
        rgb = color_to_rgb(color, workbook)
        if rgb is None:
            unresolved_colored_style = True
        elif rgb_is_blue(rgb):
            return 6
    if unresolved_colored_style:
        return None
    return 12


def header_matches(key: str, normalized: str) -> bool:
    exact: dict[str, set[str]] = {
        "name": {"姓名", "名字", "name"},
        "sim_h1": {"上半年模拟机"},
        "sim_h2": {"下半年模拟机"},
        "theory_h1": {"上半年理论复训"},
        "theory_h2": {"下半年理论复训"},
        "dangerous": {"危险品"},
        "relta": {"relta"},
        "english_level": {"英语等级"},
        "chinese_level": {"汉语等级"},
        "chinese_expiry": {"有效期"},
        "license_number": {"执照编号"},
        "id_card": {"身份证", "身份证号", "身份证号码"},
    }
    if key == "medical":
        return normalized.startswith("体检合格证")
    return normalized in exact[key]


TEST_KEYS = (
    "name",
    "sim_h1",
    "sim_h2",
    "theory_h1",
    "theory_h2",
    "dangerous",
    "medical",
    "relta",
    "english_level",
    "chinese_level",
    "chinese_expiry",
)

OPTIONAL_SENSITIVE_TEST_KEYS = ("license_number", "id_card")

CONTACT_NAME_HEADERS = {"姓名", "名字", "联系人", "name"}
CONTACT_PHONE_PRIORITY = (
    "手机号",
    "手机号码",
    "联系电话",
    "手机",
    "电话",
    "mobile",
    "phone",
)


def find_test_header(workbook: Any) -> HeaderLocation:
    candidates: list[HeaderLocation] = []
    for worksheet in workbook.worksheets:
        if getattr(worksheet, "sheet_state", "visible") != "visible":
            continue
        for row_number in range(1, min(worksheet.max_row, 30) + 1):
            values = {
                column: compact_header(worksheet.cell(row_number, column).value)
                for column in range(1, min(worksheet.max_column, 200) + 1)
            }
            columns: dict[str, int] = {}
            for key in TEST_KEYS:
                matches = [column for column, value in values.items() if header_matches(key, value)]
                if len(matches) == 1:
                    columns[key] = matches[0]
            if len(columns) == len(TEST_KEYS):
                for key in OPTIONAL_SENSITIVE_TEST_KEYS:
                    matches = [
                        column for column, value in values.items() if header_matches(key, value)
                    ]
                    if len(matches) == 1:
                        columns[key] = matches[0]
                candidates.append(HeaderLocation(worksheet.title, row_number, columns))
    if len(candidates) != 1:
        raise ConversionError(
            f"测试工作簿应且只能识别到一个完整表头，实际识别到 {len(candidates)} 个。"
        )
    return candidates[0]


def find_contact_header(workbook: Any) -> HeaderLocation:
    candidates: list[HeaderLocation] = []
    normalized_name_headers = {compact_header(value) for value in CONTACT_NAME_HEADERS}
    normalized_phone_priority = [compact_header(value) for value in CONTACT_PHONE_PRIORITY]
    for worksheet in workbook.worksheets:
        if getattr(worksheet, "sheet_state", "visible") != "visible":
            continue
        for row_number in range(1, min(worksheet.max_row, 30) + 1):
            values = {
                column: compact_header(worksheet.cell(row_number, column).value)
                for column in range(1, min(worksheet.max_column, 200) + 1)
            }
            name_columns = [
                column for column, value in values.items() if value in normalized_name_headers
            ]
            phone_column = None
            for alias in normalized_phone_priority:
                matches = [column for column, value in values.items() if value == alias]
                if matches:
                    phone_column = matches[0]
                    break
            if len(name_columns) == 1 and phone_column is not None:
                candidates.append(
                    HeaderLocation(
                        worksheet.title,
                        row_number,
                        {"name": name_columns[0], "phone": phone_column},
                    )
                )
    if len(candidates) != 1:
        raise ConversionError(
            f"通讯录应且只能识别到一个姓名/手机号表头，实际识别到 {len(candidates)} 个。"
        )
    return candidates[0]


def load_source_rows(value_workbook: Any, raw_workbook: Any) -> tuple[list[SourceRow], HeaderLocation]:
    location = find_test_header(value_workbook)
    values_sheet = value_workbook[location.sheet_name]
    raw_sheet = raw_workbook[location.sheet_name]
    rows: list[SourceRow] = []
    for row_number in range(location.row + 1, values_sheet.max_row + 1):
        name_value = values_sheet.cell(row_number, location.columns["name"]).value
        normalized_name = normalize_name(name_value)
        if not normalized_name or compact_header(name_value) == "姓名":
            continue
        raw_values = {
            key: values_sheet.cell(row_number, column).value
            for key, column in location.columns.items()
        }
        raw_cells = {
            key: raw_sheet.cell(row_number, column) for key, column in location.columns.items()
        }
        rows.append(SourceRow(row_number, normalized_name, raw_values, raw_cells))
    if not rows:
        raise ConversionError("测试工作簿没有识别到有效人员行。")
    return rows, location


def load_contact_index(workbook: Any) -> tuple[dict[str, set[str]], HeaderLocation]:
    location = find_contact_header(workbook)
    worksheet = workbook[location.sheet_name]
    index: dict[str, set[str]] = {}
    for row_number in range(location.row + 1, worksheet.max_row + 1):
        normalized_name = normalize_name(worksheet.cell(row_number, location.columns["name"]).value)
        if not normalized_name:
            continue
        raw_phone = worksheet.cell(row_number, location.columns["phone"]).value
        phone = normalize_phone(raw_phone)
        if phone:
            index.setdefault(normalized_name, set()).add(phone)
        else:
            index.setdefault(normalized_name, set())
    return index, location


def link_people(
    source_rows: Sequence[SourceRow], contact_index: dict[str, set[str]], key: bytes
) -> list[PersonRow]:
    unresolved_rows: list[int] = []
    people: list[PersonRow] = []
    person_keys: set[str] = set()
    tokens: set[str] = set()
    for source in source_rows:
        phones = contact_index.get(source.normalized_name, set())
        if len(phones) != 1:
            unresolved_rows.append(source.row_number)
            continue
        phone = next(iter(phones))
        canonical_key = f"{source.normalized_name}\x1f{phone}"
        token = person_token(key, canonical_key)
        if canonical_key in person_keys or token in tokens:
            unresolved_rows.append(source.row_number)
            continue
        person_keys.add(canonical_key)
        tokens.add(token)
        people.append(
            PersonRow(
                source.row_number,
                canonical_key,
                token,
                source.normalized_name,
                phone,
                source,
            )
        )
    if unresolved_rows:
        row_text = "、".join(str(row) for row in unresolved_rows[:30])
        suffix = "……" if len(unresolved_rows) > 30 else ""
        raise ConversionError(
            f"以下测试表行无法唯一匹配一个有效手机号，或人员键发生重复：{row_text}{suffix}。未生成输出文件。"
        )
    return people


def latest_date(
    source: SourceRow,
    keys: Sequence[str],
    excel_epoch: Any,
) -> tuple[date | None, bool, bool]:
    parsed: list[date] = []
    any_present = False
    any_invalid = False
    for key in keys:
        value = source.raw_values[key]
        raw_cell = source.raw_cells[key]
        if not is_blank(value) or getattr(raw_cell, "data_type", None) == "f":
            any_present = True
        item, invalid = parse_cell_date(value, raw_cell, excel_epoch)
        any_invalid = any_invalid or invalid
        if item:
            parsed.append(item)
    if any_invalid:
        return None, any_present, True
    return (max(parsed) if parsed else None), any_present, False


def set_qualification(
    row: dict[str, str],
    prefix: str,
    issue: date,
    expiry: date,
    level: str,
    training: date | None = None,
) -> None:
    row[f"{prefix}.issueDate"] = iso_date(issue)
    row[f"{prefix}.trainingDate"] = iso_date(training) if training else ""
    row[f"{prefix}.expiryDate"] = iso_date(expiry)
    row[f"{prefix}.levelOrParameter"] = level


def convert_person(
    person: PersonRow,
    mock_mobile: str,
    excel_epoch: Any,
    raw_workbook: Any,
    warn: Callable[[int, str], None],
) -> dict[str, str]:
    source = person.source
    row = {header: "" for header in EXPECTED_HEADERS}
    row.update(
        {
            "employeeNumber": f"MOCK-{person.token}",
            "displayName": f"飞行员-{person.token}",
            "mobile": mock_mobile,
        }
    )

    simulator_date, simulator_present, simulator_invalid = latest_date(
        source, ("sim_h1", "sim_h2"), excel_epoch
    )
    if simulator_date and not simulator_invalid:
        set_qualification(
            row,
            "simulator-recurrent-training",
            simulator_date,
            add_months(simulator_date, 6),
            "合格",
            simulator_date,
        )
    elif simulator_present:
        warn(source.row_number, "模拟机")

    theory_date, theory_present, theory_invalid = latest_date(
        source, ("theory_h1", "theory_h2"), excel_epoch
    )
    if theory_date and not theory_invalid:
        set_qualification(
            row,
            "annual-recurrent-training",
            theory_date,
            add_months(theory_date, 12),
            "合格",
            theory_date,
        )
    elif theory_present:
        warn(source.row_number, "理论复训")

    dangerous_date, dangerous_invalid = parse_cell_date(
        source.raw_values["dangerous"], source.raw_cells["dangerous"], excel_epoch
    )
    if dangerous_date and not dangerous_invalid:
        set_qualification(
            row,
            "dangerous-goods-training",
            dangerous_date,
            add_months(dangerous_date, 24),
            "合格",
        )
    elif not is_blank(source.raw_values["dangerous"]):
        warn(source.row_number, "危险品")

    medical_date, medical_invalid = parse_cell_date(
        source.raw_values["medical"], source.raw_cells["medical"], excel_epoch
    )
    if medical_date and not medical_invalid:
        cycle = medical_cycle_months(source.raw_cells["medical"], raw_workbook)
        if cycle is None:
            warn(source.row_number, "体检颜色")
        else:
            set_qualification(
                row,
                "medical-certificate",
                medical_date,
                add_months(medical_date, cycle),
                "半年体检" if cycle == 6 else "年度体检",
            )
    elif not is_blank(source.raw_values["medical"]):
        warn(source.row_number, "体检日期")

    english_date, english_invalid = parse_cell_date(
        source.raw_values["relta"], source.raw_cells["relta"], excel_epoch
    )
    english_level = parse_level(
        source.raw_values["english_level"], source.raw_cells["english_level"]
    )
    english_present = not is_blank(source.raw_values["relta"]) or not is_blank(
        source.raw_values["english_level"]
    )
    if english_date and english_level and not english_invalid:
        set_qualification(
            row,
            "icao-english-endorsement",
            english_date,
            add_months(english_date, 36),
            english_level,
        )
    elif english_present:
        warn(source.row_number, "英语语言资质")

    chinese_expiry, chinese_invalid = parse_cell_date(
        source.raw_values["chinese_expiry"],
        source.raw_cells["chinese_expiry"],
        excel_epoch,
    )
    chinese_level = parse_level(
        source.raw_values["chinese_level"], source.raw_cells["chinese_level"]
    )
    chinese_present = not is_blank(source.raw_values["chinese_expiry"]) or not is_blank(
        source.raw_values["chinese_level"]
    )
    if chinese_expiry and chinese_level and not chinese_invalid:
        set_qualification(
            row,
            "chinese-language-assessment",
            add_months(chinese_expiry, -36),
            chinese_expiry,
            chinese_level,
        )
    elif chinese_present:
        warn(source.row_number, "汉语语言资质")

    return row


def read_and_validate_template(path: Path) -> list[str]:
    if not path.is_file():
        raise ConversionError(f"找不到当前 CrewQual 模板：{path}")
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = [row for row in csv.reader(handle) if any(cell.strip() for cell in row)]
    if not rows:
        raise ConversionError("CrewQual 模板为空。")
    if rows[0] != EXPECTED_HEADERS:
        raise ConversionError("CrewQual 模板表头不是当前预期的 31 列，拒绝生成不兼容文件。")
    return rows[0]


def render_csv(headers: Sequence[str], rows: Sequence[dict[str, str]]) -> bytes:
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow(headers)
    for row in rows:
        writer.writerow([row[header] for header in headers])
    return ("\ufeff" + buffer.getvalue()).encode("utf-8")


def validate_output(
    headers: Sequence[str],
    rows: Sequence[dict[str, str]],
    people: Sequence[PersonRow],
    key: bytes,
) -> None:
    if list(headers) != EXPECTED_HEADERS or len(headers) != 31:
        raise ConversionError("输出表头自检失败。")
    if len(rows) != len(people):
        raise ConversionError("输出人数与源表有效人数不一致。")

    employee_numbers = [row["employeeNumber"] for row in rows]
    display_names = [row["displayName"] for row in rows]
    mobiles = [row["mobile"] for row in rows]
    if any(len(set(values)) != len(values) for values in (employee_numbers, display_names, mobiles)):
        raise ConversionError("mock 员工号、姓名或手机号发生碰撞。")
    if not all(re.fullmatch(r"MOCK-[A-Z2-7]{16}", value) for value in employee_numbers):
        raise ConversionError("mock 员工号格式自检失败。")
    if not all(re.fullmatch(r"100\d{8}", value) for value in mobiles):
        raise ConversionError("mock 手机号格式自检失败。")

    sensitive_names = {person.normalized_name for person in people}
    sensitive_phones = {person.real_phone for person in people}
    sensitive_identifiers = {
        normalize_text(person.source.raw_values.get(key))
        for person in people
        for key in OPTIONAL_SENSITIVE_TEST_KEYS
        if normalize_text(person.source.raw_values.get(key))
    }
    for row in rows:
        normalized_cells = [normalize_name(value) for value in row.values() if value]
        raw_cells = {normalize_text(value) for value in row.values() if value}
        if any(name and name in normalized_cells for name in sensitive_names):
            raise ConversionError("输出仍包含真实姓名。")
        if any(phone in row.values() for phone in sensitive_phones):
            raise ConversionError("输出仍包含真实手机号。")
        if any(identifier in raw_cells for identifier in sensitive_identifiers):
            raise ConversionError("输出仍包含真实身份证或执照编号。")
        if key.hex() in raw_cells:
            raise ConversionError("输出意外包含 HMAC 密钥。")
        for prefix, training_required in QUALIFICATION_GROUPS.items():
            issue = row[f"{prefix}.issueDate"]
            training = row[f"{prefix}.trainingDate"]
            expiry = row[f"{prefix}.expiryDate"]
            level = row[f"{prefix}.levelOrParameter"]
            values = (issue, training, expiry, level)
            if not any(values):
                continue
            if not issue or not expiry or not level or (training_required and not training):
                raise ConversionError(f"资质 {prefix} 存在半组数据。")
            issue_date = parse_date_value(issue)
            expiry_date = parse_date_value(expiry)
            if not issue_date or not expiry_date or expiry_date < issue_date:
                raise ConversionError(f"资质 {prefix} 的日期自检失败。")

    alternate_key = hashlib.sha256(key + b"CrewQual alternate-key self-test").digest()
    alternate_mobiles = assign_mock_mobiles(alternate_key, people)
    for person, row in zip(people, rows, strict=True):
        alternate_token = person_token(alternate_key, person.person_key)
        if (
            row["employeeNumber"] == f"MOCK-{alternate_token}"
            or row["displayName"] == f"飞行员-{alternate_token}"
            or row["mobile"] == alternate_mobiles[person.token]
        ):
            raise ConversionError("更换密钥后假名未变化，自检失败。")
    first = render_csv(headers, rows)
    if first != render_csv(headers, rows):
        raise ConversionError("相同输入的 CSV 字节结果不稳定。")


def qualification_counts(rows: Sequence[dict[str, str]]) -> dict[str, int]:
    return {
        prefix: sum(bool(row[f"{prefix}.issueDate"]) for row in rows)
        for prefix in QUALIFICATION_GROUPS
    }


def discover_contacts(source_path: Path, explicit: Path | None) -> Path:
    if explicit:
        return explicit
    candidates = sorted(
        path
        for path in source_path.parent.glob("*通讯录*.xlsx")
        if not path.name.startswith("~$") and path.resolve() != source_path.resolve()
    )
    if len(candidates) != 1:
        raise ConversionError(
            f"源文件目录中应且只能有一个文件名含“通讯录”的 XLSX，实际找到 {len(candidates)} 个；可用 --contacts 指定。"
        )
    return candidates[0]


def convert(
    source_path: Path,
    contacts_path: Path,
    template_path: Path,
    output_path: Path,
    overwrite: bool,
) -> ConversionSummary:
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise ConversionError("缺少 openpyxl；请先运行 setup_windows.bat。") from error

    if not source_path.is_file():
        raise ConversionError(f"找不到测试工作簿：{source_path}")
    if not contacts_path.is_file():
        raise ConversionError(f"找不到通讯录工作簿：{contacts_path}")
    protected_paths = {source_path.resolve(), contacts_path.resolve(), template_path.resolve()}
    if output_path.resolve() in protected_paths:
        raise ConversionError("输出路径不能与测试工作簿、通讯录或当前模板相同。")
    if output_path.exists() and not overwrite:
        raise ConversionError(f"输出文件已存在：{output_path}；确认后使用 --overwrite 覆盖。")
    headers = read_and_validate_template(template_path)
    key = load_hmac_key()

    try:
        source_values = load_workbook(source_path, data_only=True, read_only=False)
        source_raw = load_workbook(source_path, data_only=False, read_only=False)
        contacts = load_workbook(contacts_path, data_only=True, read_only=True)
    except Exception as error:
        raise ConversionError(f"无法读取 XLSX：{type(error).__name__}。") from error

    try:
        source_rows, _ = load_source_rows(source_values, source_raw)
        contact_index, _ = load_contact_index(contacts)
        people = link_people(source_rows, contact_index, key)
        mock_mobiles = assign_mock_mobiles(key, people)
        warning_items: list[tuple[int, str]] = []

        def warn(row_number: int, category: str) -> None:
            warning_items.append((row_number, category))

        rows = [
            convert_person(
                person,
                mock_mobiles[person.token],
                source_values.epoch,
                source_raw,
                warn,
            )
            for person in people
        ]
        validate_output(headers, rows, people, key)
        payload = render_csv(headers, rows)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        temporary = output_path.with_name(f".{output_path.name}.tmp")
        try:
            temporary.write_bytes(payload)
            os.replace(temporary, output_path)
        finally:
            temporary.unlink(missing_ok=True)
    finally:
        source_values.close()
        source_raw.close()
        contacts.close()

    for row_number, category in warning_items:
        print(f"[警告] 测试表第 {row_number} 行的“{category}”无法可靠转换，该资质整组留空。")
    counts = qualification_counts(rows)
    return ConversionSummary(len(source_rows), len(people), counts, len(warning_items))


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="将测试.xlsx和通讯录转换为稳定假名化的 CrewQual CSV 草稿。"
    )
    parser.add_argument("--source", type=Path, default=Path(DEFAULT_SOURCE), help="测试 XLSX")
    parser.add_argument("--contacts", type=Path, help="通讯录 XLSX；默认自动寻找 *通讯录*.xlsx")
    parser.add_argument(
        "--template", type=Path, default=Path(DEFAULT_TEMPLATE), help="CrewQual 当前 CSV 模板"
    )
    parser.add_argument("--output", type=Path, default=Path(DEFAULT_OUTPUT), help="输出 CSV")
    parser.add_argument("--overwrite", action="store_true", help="覆盖已存在的输出 CSV")
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        source = args.source.resolve()
        contacts = discover_contacts(source, args.contacts.resolve() if args.contacts else None)
        template = args.template.resolve()
        output = args.output.resolve()
        summary = convert(source, contacts, template, output, args.overwrite)
    except ConversionError as error:
        print(f"[错误] {error}", file=sys.stderr)
        return 2
    except Exception as error:
        print(f"[错误] 转换意外失败：{type(error).__name__}。未输出敏感数据。", file=sys.stderr)
        return 3

    print(f"[完成] 已生成：{output}")
    print(
        f"[验收] 源表人员 {summary.source_rows}，唯一匹配手机号 {summary.matched_phones}，"
        f"留空资质 {summary.suppressed_qualifications}。"
    )
    labels = {
        "medical-certificate": "体检",
        "annual-recurrent-training": "理论复训",
        "dangerous-goods-training": "危险品",
        "icao-english-endorsement": "英语",
        "chinese-language-assessment": "汉语",
        "simulator-recurrent-training": "模拟机",
    }
    print(
        "[验收] 资质写入："
        + "、".join(f"{labels[prefix]} {count}" for prefix, count in summary.qualification_counts.items())
    )
    print("[提示] 该文件保留原始业务日期，属于假名化 mock 数据，不是匿名数据。")
    print("[提示] 机型、职务、单位代码、人员级别仍为空，补齐前不能正式导入 CrewQual。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
